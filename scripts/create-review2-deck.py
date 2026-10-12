"""Create an editable three-slide Review 2 deck using the supplied SRM template."""
import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
if len(sys.argv) != 2:
    raise SystemExit('Usage: python scripts/create-review2-deck.py <review-template.pptx>')
SOURCE = Path(sys.argv[1]).expanduser().resolve()
report = json.loads((ROOT / 'outputs/review2-evidence.json').read_text(encoding='utf-8'))
p = Presentation(SOURCE)
members = [row.cells[1].text.strip() for slide in p.slides for shape in slide.shapes if shape.has_table for row in shape.table.rows]
PURPLE = '5900C2'
BLACK = '161616'
GRAY = '555555'

def text(slide, value, x, y, w, h, size=20, bold=False, color=BLACK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.03)
    frame.margin_top = frame.margin_bottom = Inches(.03)
    for index, line in enumerate(value.split('\n')):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = line
        para.space_after = Pt(7)
        para.font.name = 'Cambria'
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = RGBColor.from_string(color)
    return box

def title(slide, value, index):
    text(slide, value, 1.12, .28, 11.0, .65, 31, True, PURPLE)
    text(slide, f'SignBridge | Review 2 | 7 October 2026', .6, 6.98, 10.7, .25, 11, color=GRAY)
    text(slide, str(index), 12.1, 6.95, .35, .3, 12, color=GRAY)

def table(slide, rows, widths, x, y, h, size=19):
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(h))
    tab = shape.table
    for i, width in enumerate(widths): tab.columns[i].width = Inches(width)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tab.cell(i,j)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string('EEE8F6' if i == 0 else 'FFFFFF')
            cell.margin_left = cell.margin_right = Inches(.1)
            cell.margin_top = cell.margin_bottom = Inches(.055)
            for para in cell.text_frame.paragraphs:
                para.font.name = 'Cambria'
                para.font.size = Pt(size)
                para.font.bold = i == 0
                para.font.color.rgb = RGBColor.from_string(PURPLE if i == 0 else BLACK)
    return shape

def node(slide, label, x, y, w=2.12):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(.85))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string('EEE8F6')
    shape.line.color.rgb = RGBColor.from_string(PURPLE)
    frame = shape.text_frame; frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.08)
    para = frame.paragraphs[0]; para.text = label; para.alignment = PP_ALIGN.CENTER
    para.font.name = 'Cambria'; para.font.size = Pt(18); para.font.bold = True
    para.font.color.rgb = RGBColor.from_string(BLACK)

for slide in p.slides:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)

s = p.slides[0]
title(s, 'GenAI justification and Review 1 feedback', 1)
text(s, 'Implemented: local English paraphrasing with Qwen2.5-1.5B-Instruct', .65, 1.18, 12, .5, 23, True)
rows = [['Concept', 'Use in SignBridge'], ['ML / DL', 'Pretrained neural model weights'], ['NLP', 'Rewrite English while preserving details'], ['Transformer', 'Attention over source and prompt context'], ['GenAI', 'Generate a paraphrase token by token']]
table(s, rows, [2.05,5.95], .65, 1.95, 2.85)
text(s, 'Why this model', 9.0, 1.95, 3.45, .42, 22, True, PURPLE)
text(s, 'Runs on the local CPU\nNo paid API credentials\nActual prompt/output evidence\nStructured JSON validation', 9.0, 2.5, 3.55, 2.4, 20)
text(s, 'Review 1: the idea was appreciated; no specific corrections requested (team-reported).', .65, 5.12, 11.8, .65, 20, True)
text(s, 'Added for Review 2: real model inference, prompt experiments and an integrated language lab.\nPretrained inference only; no model fine-tuning or autonomous agent workflow.', .65, 5.85, 11.8, .92, 18)
s.notes_slide.notes_text_frame.text = '''1:45 total, including Review 1 feedback.
SignBridge aims to support text-to-sign communication. Our new implemented GenAI stage generates clearer English using a pretrained Qwen Transformer. Machine learning supplies learned weights; deep learning supplies neural layers; NLP is the English rewriting task; attention processes source and prompt context. The model predicts output tokens rather than selecting a stored paraphrase. We use a quantized model locally because this prototype has no paid API requirement. Every run exposes its exact prompt and measured timing. We did not train or fine-tune model weights. The Review 1 feedback, as reported by our team, was positive with no requested corrections. For this review we added an actual model, the language lab and experiments. We still need a separate trained ISL motion model for unrestricted signing.
Sources: https://huggingface.co/Qwen/Qwen2.5-1.5B-Instruct ; https://huggingface.co/Qwen/Qwen2.5-1.5B-Instruct-GGUF ; https://github.com/ggml-org/llama.cpp
All team members must attend: ''' + ', '.join(members)

s = p.slides[1]
title(s, 'Generated content, objectives and pipeline', 2)
labels = ['English\ninput', 'Prompt +\nQwen model', 'English\nparaphrase', 'Checks +\nconfirmation', 'ISL asset gate\n+ 3D playback']
for i,label in enumerate(labels):
    x=.65+i*2.43; node(s,label,x,1.35)
    if i<4:
        arrow=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,Inches(x+2.17),Inches(1.65),Inches(.22),Inches(.16))
        arrow.fill.solid();arrow.fill.fore_color.rgb=RGBColor.from_string(PURPLE);arrow.line.fill.background()
text(s, 'English generation works for new inputs. Avatar playback uses available reference assets.\nUnsupported ISL sentences remain blocked; no signing-model weights are installed.', .65, 2.43, 11.9, .88, 19)
text(s, 'Objective: preserve names, time, negation and requests', .65, 3.46, 12, .45, 23, True, PURPLE)
table(s, [['Actual scenario', 'Observed issue', 'Constrained output'], ['Certificate request', 'Basic: "valid certificate"', '"bonafide certificate"'], ['Library closing time', 'Examples: "today" omitted', 'Retains "today"'], ['Appointment contrast', 'Both time needs preserved', 'Today / tomorrow retained']], [3.25,4.15,4.5], .65, 4.08, 1.65, 18)
text(s, 'Analysis: inspect omissions and substitutions; literal checks do not prove meaning.\nThe pipeline diagram was refined. ISL generation and expert evaluation remain outstanding.', .65, 6.0, 12, .8, 18)
s.notes_slide.notes_text_frame.text = '''1:45 total.
This is the updated implemented pipeline. English enters a validated API, goes to the selected prompt and pretrained Transformer, and returns an English paraphrase. We then run preservation checks and ask the user to inspect meaning before explicitly handing it to the studio. The existing sign provider and asset gate decide whether an installed reference motion can play. This is not yet complete sentence-level ISL translation. Our objective for this implemented stage is to preserve the message during simplification. We tested a campus question, negation and time, an unfamiliar place, a long paragraph, and a numbered conditional request. A basic prompt changed bonafide to valid in one actual result, while the constrained prompt retained bonafide. The example prompt omitted today in the library question. These show why names and time need explicit instructions. All outputs and prompts are available on the review page. Literal matching is only a diagnostic, and independent human meaning review remains needed. The diagram refinement is complete; we did not fine-tune model weights.
Evidence: outputs/review2-evidence.json. Table quotations are excerpts from actual recorded responses, not complete translations.'''

s = p.slides[2]
title(s, 'Prompt engineering: measured use cases', 3)
text(s, '5 scenarios x 3 prompts = 15 actual local-model generations', .65, 1.18, 12, .45, 23, True)
rows = [['Prompt strategy','Successful runs','Terms retained','Mean time']]
names = {'zero-shot':'Basic instruction','few-shot':'With worked examples','constrained':'Preservation rules + examples'}
for summary in report['summaries']:
    rows.append([names[summary['strategy']],f"{summary['successfulRuns']}/{summary['totalRuns']}",f"{summary['literalTermsFound']}/{summary['literalTermsTotal']}",f"{summary['meanLatencyMs']/1000:.1f}s"])
table(s,rows,[4.4,2.35,2.5,2.65],.65,1.98,1.95,19)
text(s, 'Prompt change', .65, 4.15, 3.0, .42, 23, True, PURPLE)
text(s, '"Keep names, place names, abbreviations, numbers, dates and times exactly.\nKeep negation, uncertainty and conditions. Keep each question as a question."', .65, 4.69, 11.9, .92, 20)
text(s, 'Use cases: campus help, unfamiliar names, paragraphs, opposite time needs and conditions.\nSame model/schema, temperature 0, seed 42. Term retention is not semantic or ISL accuracy.', .65, 5.91, 11.9, .86, 18)
s.notes_slide.notes_text_frame.text = '''1:30 total. This completes the five-minute presentation.
We compared three prompt strategies on the same five inputs using the same model, output schema, temperature zero and seed 42. The basic prompt directly requests paraphrasing. The second supplies worked examples that differ from the evaluation inputs. The constrained version adds explicit preservation rules. The table shows actual measured runs, not estimated performance. The constrained version retained all seventeen tracked terms in these examples, while the other versions lost some. This is not a percentage of ISL accuracy or a general guarantee. Timings include the model call and vary with hardware, cache reuse and warm-up. The key prompt improvement is making preservation of names, numbers, negation, uncertainty and question form explicit. Demonstrate the library question in the language lab and open the exact prompt and result if time permits. We have a working integrated web prototype with local generation, persistence and avatar controls, but arbitrary-text ISL generation and expert review remain outstanding; our checklist should not be presented as seventy percent of the original unrestricted translation objective.
Evidence recorded: ''' + report['generatedAt'] + '\n' + report['protocol']

target = ROOT / 'outputs/SignBridge-Review2.pptx'
p.save(target)
print(f'Saved {target} with {len(p.slides)} slides. Notes: 105 + 105 + 90 seconds.')
